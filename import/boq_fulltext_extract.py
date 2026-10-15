#!/usr/bin/env python3
"""Re-read the 58 sent BOQ workbooks that fed the vault's rates sweep and emit every priced
line with its FULL description (the sweep clipped them at 90 chars). Same FINAL-sheet
parser and geometry as the BOQ assistant's boq_rates_sweep.py.

Runs on a machine that can see the quotations share (Goshi):
  python3 boq_fulltext_extract.py "<root>" sweep_files.txt > fulltext.json
Writes a log beside the output. Output: list of {file, ref, date, section, desc, qty, unit, rate, cost}.
"""
import json, re, sys, time
from pathlib import Path
import openpyxl

C_NO, C_DESC, C_QTY, C_UNIT, C_RATE, C_AMT, C_COST = 1, 3, 6, 7, 8, 9, 10
QUOTE_REF = re.compile(r"BBC/\d{6}/\d{3,6}(?:_rev\d+)?", re.I)
SPECIAL = [("by_client", re.compile(r"(?i)by\s+(the\s+)?client|by\s+others")),
           ("excluded", re.compile(r"(?i)to be confirmed and quoted separately|not included|excluded")),
           ("rate_only", re.compile(r"(?i)rate\s*only"))]
SUBTOTAL = re.compile(r"(?i)^(sub-?total|total of)\b"); TOTALS = re.compile(r"(?i)^(total\s*$|vat\b|grand total|summary)")

def num(v):
    if isinstance(v, (int, float)): return float(v)
    if isinstance(v, str):
        try: return float(v.replace(",", "").replace("AED", "").strip())
        except ValueError: return None
    return None
def cellv(vals, c): return vals[c - 1] if len(vals) >= c else None
def desc_of(vals):
    for c in (C_DESC, C_NO, C_DESC + 1, 2):
        v = cellv(vals, c)
        if isinstance(v, str) and len(v.strip()) >= 5: return v.strip()
    return None
def classify(vals):
    d = desc_of(vals) or ""; q = num(cellv(vals, C_QTY))
    price = " ".join(str(cellv(vals, c)) for c in (C_RATE, C_AMT) if isinstance(cellv(vals, c), str))
    for f, p in SPECIAL:
        if p.search(price): return f
    if SUBTOTAL.match(d): return "subtotal"
    if TOTALS.match(d): return "totals"
    if q is None:
        for f, p in SPECIAL:
            if p.search(d): return f
        return "text"
    return "priced" if d else "text"

def parse(path):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        name = next((s for s in wb.sheetnames if s.strip().upper() == "FINAL"), None)
        if name is None: return None, [], "no FINAL sheet"
        rows = [list(v) for v in wb[name].iter_rows(values_only=True)]
    finally: wb.close()
    ref = ""
    for vals in rows[:40]:
        for v in vals:
            if isinstance(v, str) and (m := QUOTE_REF.search(v)): ref = m.group(0); break
        if ref: break
    if not ref and (m := QUOTE_REF.search(path.name)): ref = m.group(0)
    section, out = "", []
    for vals in rows:
        d = desc_of(vals)
        if d is None: continue
        k = classify(vals)
        if k == "text":
            if re.match(r"^([IVX]+|[A-Z])[\.\)]\s+\S", d): section = d
            continue
        if k in ("subtotal", "totals", "by_client", "excluded"): continue
        rate = num(cellv(vals, C_RATE))
        if rate is None or rate <= 0: continue
        out.append(dict(section=section, desc=d, qty=num(cellv(vals, C_QTY)), unit=cellv(vals, C_UNIT), rate=rate, cost=num(cellv(vals, C_COST))))
    return ref, out, None

def main():
    root = Path(sys.argv[1]); files = [l.strip() for l in open(sys.argv[2], encoding="utf-8") if l.strip()]
    log = open(Path(sys.argv[3]) if len(sys.argv) > 3 else "fulltext-log.txt", "w", encoding="utf-8")
    say = lambda m: (log.write(m + "\n"), log.flush(), sys.stderr.write(m + "\n"))
    say(f"=== run {time.strftime('%Y-%m-%d %H:%M:%S')} root={root} files={len(files)}")
    all_rows = []
    for rel in files:
        p = root / rel.replace("\\", "/")
        if not p.exists(): say(f"MISSING {rel}"); continue
        try:
            ref, rows, err = parse(p)
        except Exception as e:
            say(f"ERROR {rel}: {e}"); continue
        if err: say(f"SKIP {rel}: {err}"); continue
        date = time.strftime("%Y-%m-%d", time.localtime(p.stat().st_mtime))
        for r in rows: all_rows.append(dict(file=rel, ref=ref, date=date, **r))
        say(f"OK {len(rows):4d} lines  {rel}")
    say(f"=== done: {len(all_rows)} priced lines")
    json.dump(all_rows, sys.stdout, ensure_ascii=False)

if __name__ == "__main__": main()
